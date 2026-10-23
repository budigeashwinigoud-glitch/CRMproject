import io
import csv
import re
from typing import List, Dict, Any, Tuple, Optional
import openpyxl
from sqlalchemy.orm import Session
from app.models.lead import Lead
from app.schemas.lead import LeadCreate
from app.schemas.import_export import ImportReport, RowError, DuplicateMatch
from app.services.lead_service import create_lead
from app.services.duplicate_service import check_lead_duplicates, normalize_email, normalize_phone


def sanitize_formula_injection(value: Any) -> Any:
    """Prepend a single quote to prevent spreadsheet formula injection."""
    if isinstance(value, str) and value and value[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{value}"
    return value


def parse_csv_file(file_bytes: bytes) -> List[Dict[str, Any]]:
    text = file_bytes.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    rows = []
    for r in reader:
        # Strip keys and values
        clean_row = {k.strip(): (v.strip() if v else "") for k, v in r.items() if k}
        rows.append(clean_row)
    return rows


def parse_excel_file(file_bytes: bytes) -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    data_rows = []
    for row in rows[1:]:
        if not any(row):
            continue
        row_dict = {}
        for h, val in zip(headers, row):
            if h:
                row_dict[h] = str(val).strip() if val is not None else ""
        data_rows.append(row_dict)
    return data_rows


def process_lead_import(
    db: Session,
    raw_rows: List[Dict[str, Any]],
    dry_run: bool = False,
    update_duplicates: bool = False,
    performed_by: str = "Pavan Kumar"
) -> ImportReport:
    total_rows = len(raw_rows)
    inserted = 0
    updated = 0
    skipped_duplicates = 0
    rejected_errors = 0
    errors: List[RowError] = []
    duplicate_warnings: List[DuplicateMatch] = []

    for idx, row in enumerate(raw_rows, start=1):
        # Map common column naming variations
        full_name = row.get("fullName") or row.get("Full Name") or row.get("name") or row.get("Contact") or ""
        company = row.get("company") or row.get("Company") or row.get("Organization") or ""
        email = row.get("email") or row.get("Email") or ""
        phone = row.get("phone") or row.get("Phone") or row.get("Mobile") or ""
        source = row.get("leadSource") or row.get("Lead Source") or row.get("source") or "Website"
        category = row.get("leadCategory") or row.get("Lead Category") or row.get("category") or "Enterprise Banking"
        priority = row.get("priority") or row.get("Priority") or "Medium"
        stage = row.get("pipelineStage") or row.get("Pipeline Stage") or row.get("stage") or "New"
        assigned = row.get("assignedEmployee") or row.get("Assigned Employee") or row.get("employee") or performed_by
        value_raw = row.get("expectedValue") or row.get("Expected Value") or row.get("value") or "0"
        notes = row.get("notes") or row.get("Notes") or ""
        website = row.get("website") or row.get("Website") or ""
        location = row.get("location") or row.get("Location") or ""
        follow_up = row.get("nextFollowUpDate") or row.get("Next Follow Up Date") or ""

        # Validate required fields
        if not company:
            errors.append(RowError(rowNumber=idx, field="company", message="Company name is required."))
            rejected_errors += 1
            continue

        if not full_name:
            errors.append(RowError(rowNumber=idx, field="fullName", message="Contact full name is required."))
            rejected_errors += 1
            continue

        if not email or "@" not in email:
            errors.append(RowError(rowNumber=idx, field="email", message=f"Invalid email address: '{email}'"))
            rejected_errors += 1
            continue

        try:
            expected_val = float(re.sub(r"[^\d.]", "", str(value_raw))) if str(value_raw).strip() else 0.0
        except Exception:
            expected_val = 0.0

        # Check duplicates
        matches = check_lead_duplicates(db, email=email, phone=phone, company=company, full_name=full_name)
        if matches:
            duplicate_warnings.extend(matches)
            if not update_duplicates:
                skipped_duplicates += 1
                continue
            else:
                # Update first matching lead
                if not dry_run:
                    match_lead = db.query(Lead).filter(Lead.id == matches[0].existingLeadId).first()
                    if match_lead:
                        match_lead.expected_value = expected_val or match_lead.expected_value
                        match_lead.priority = priority or match_lead.priority
                        match_lead.notes = f"{match_lead.notes}\n[Import Update]: {notes}" if notes else match_lead.notes
                        db.commit()
                updated += 1
                continue

        # If clean lead, insert
        if not dry_run:
            lead_in = LeadCreate(
                fullName=full_name,
                company=company,
                email=email,
                phone=phone,
                leadSource=source,
                leadCategory=category,
                priority=priority,
                pipelineStage=stage,
                assignedEmployee=assigned,
                expectedValue=expected_val,
                notes=notes,
                website=website,
                location=location,
                nextFollowUpDate=follow_up
            )
            create_lead(db, lead_in, performed_by=performed_by)

        inserted += 1

    return ImportReport(
        totalRows=total_rows,
        inserted=inserted,
        updated=updated,
        skippedDuplicates=skipped_duplicates,
        rejectedErrors=rejected_errors,
        errors=errors,
        duplicateWarnings=duplicate_warnings
    )


def export_leads_to_csv(leads: List[Lead]) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    headers = [
        "ID", "Full Name", "Company", "Email", "Phone", "Lead Source",
        "Lead Category", "Priority", "Pipeline Stage", "Assigned Employee",
        "Expected Value (USD)", "Next Follow Up Date", "Created Date", "Location", "Notes"
    ]
    writer.writerow(headers)

    for l in leads:
        writer.writerow([
            sanitize_formula_injection(l.id),
            sanitize_formula_injection(l.full_name),
            sanitize_formula_injection(l.company),
            sanitize_formula_injection(l.email),
            sanitize_formula_injection(l.phone),
            sanitize_formula_injection(l.lead_source),
            sanitize_formula_injection(l.lead_category),
            sanitize_formula_injection(l.priority),
            sanitize_formula_injection(l.pipeline_stage),
            sanitize_formula_injection(l.assigned_employee),
            l.expected_value or 0.0,
            sanitize_formula_injection(l.next_follow_up_date),
            sanitize_formula_injection(l.created_date),
            sanitize_formula_injection(l.location),
            sanitize_formula_injection(l.notes)
        ])

    return output.getvalue()


def export_leads_to_excel(leads: List[Lead]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"

    headers = [
        "ID", "Full Name", "Company", "Email", "Phone", "Lead Source",
        "Lead Category", "Priority", "Pipeline Stage", "Assigned Employee",
        "Expected Value (USD)", "Next Follow Up Date", "Created Date", "Location", "Notes"
    ]
    ws.append(headers)

    for l in leads:
        ws.append([
            sanitize_formula_injection(l.id),
            sanitize_formula_injection(l.full_name),
            sanitize_formula_injection(l.company),
            sanitize_formula_injection(l.email),
            sanitize_formula_injection(l.phone),
            sanitize_formula_injection(l.lead_source),
            sanitize_formula_injection(l.lead_category),
            sanitize_formula_injection(l.priority),
            sanitize_formula_injection(l.pipeline_stage),
            sanitize_formula_injection(l.assigned_employee),
            l.expected_value or 0.0,
            sanitize_formula_injection(l.next_follow_up_date),
            sanitize_formula_injection(l.created_date),
            sanitize_formula_injection(l.location),
            sanitize_formula_injection(l.notes)
        ])

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
