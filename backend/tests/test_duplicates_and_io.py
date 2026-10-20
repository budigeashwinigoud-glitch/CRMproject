import io
import openpyxl


def test_duplicate_detection_endpoint(client, admin_token):
    headers = {"Authorization": f"Bearer {admin_token}"}
    # Test existing email match
    payload = {
        "email": "rodrigo.sanchez@nubank.com",
        "company": "Another Bank",
        "fullName": "Rodrigo Silva"
    }
    res = client.post("/api/leads/duplicates/check", json=payload, headers=headers)
    assert res.status_code == 200
    data = res.json()
    assert data["isDuplicate"] is True
    assert len(data["potentialMatches"]) >= 1
    assert data["potentialMatches"][0]["confidence"] == "High"

    # Test fresh non-duplicate lead
    clean_payload = {
        "email": "unique.founder@novelstart.com",
        "company": "Novel FinTech Ltd",
        "fullName": "Mariana Nova"
    }
    clean_res = client.post("/api/leads/duplicates/check", json=clean_payload, headers=headers)
    assert clean_res.status_code == 200
    assert clean_res.json()["isDuplicate"] is False


def test_csv_import_and_validation(client, admin_token):
    headers = {"Authorization": f"Bearer {admin_token}"}
    csv_content = """fullName,company,email,phone,leadSource,priority,expectedValue
Lucas Pires,Stone Pagamentos,lucas.pires@stone.co,+55 11 91234-5678,Partner,High,88000
Invalid Row,,bad-email,,Website,Low,10000
"""
    files = {"file": ("test_import.csv", io.BytesIO(csv_content.encode("utf-8")), "text/csv")}
    res = client.post("/api/import/leads", files=files, headers=headers)
    assert res.status_code == 200
    report = res.json()
    assert report["totalRows"] == 2
    assert report["inserted"] == 1
    assert report["rejectedErrors"] == 1
    assert len(report["errors"]) == 1
    assert report["errors"][0]["rowNumber"] == 2


def test_csv_export_with_formula_sanitization(client, admin_token):
    headers = {"Authorization": f"Bearer {admin_token}"}
    res = client.get("/api/export/leads?format=csv", headers=headers)
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/csv")
    csv_text = res.text
    assert "ID,Full Name,Company,Email" in csv_text
    assert "Nubank Brasil" in csv_text


def test_excel_export(client, admin_token):
    headers = {"Authorization": f"Bearer {admin_token}"}
    res = client.get("/api/export/leads?format=xlsx", headers=headers)
    assert res.status_code == 200
    assert "openxmlformats" in res.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Leads" in wb.sheetnames
    sheet = wb["Leads"]
    assert sheet.max_row >= 10
